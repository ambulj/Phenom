"""
Unit and Integration Tests for Phase 5: Modes, Scheduler, and Integrations
Verifies:
1. Scheduler fires on time and survives a database restart.
2. Live data API failure falls back to scraping and is labelled ('source': 'scraped').
3. Financial watchlist price alert triggers once and never duplicates.
4. Hard prohibitions enforced in code: Email sending, calendar edit/delete, trading, advice.
5. All 4 document generators create valid files that open and parse correctly (.docx, .xlsx, .pptx, .pdf).
6. Assistant mode switching (Work, Study, Chill) modifies tone and notification rules.
7. Tech news briefing aggregation.
"""

import os
import pytest
import asyncio
from datetime import datetime, timezone, timedelta
import docx
import openpyxl
import pptx
import pypdf

from services.backend.memory.db import EncryptedDatabaseManager
from services.backend.modes.mode_manager import ModeManager, AssistantMode
from services.backend.scheduler.cron import BackgroundScheduler
from services.backend.integrations.finance import FinancialEngine, TradingProhibitedError
from services.backend.integrations.google_workspace import (
    GoogleWorkspaceClient, EmailSendingProhibitedError, CalendarModificationProhibitedError
)
from services.backend.integrations.live_data import LiveDataPipeline
from services.backend.integrations.news_briefing import TechNewsBriefing
from services.backend.integrations.research import ResearchEngine
from services.backend.integrations.doc_generator import DocumentGenerator
from services.backend.safety.registry import registry
import services.backend.tools.integration_tools

@pytest.fixture
def temp_db(tmp_path):
    db_file = str(tmp_path / "test_phase5.db")
    mgr = EncryptedDatabaseManager(db_path=db_file, encryption_key="0123456789abcdef0123456789abcdef")
    return mgr

@pytest.mark.asyncio
async def test_scheduler_fires_on_time_and_survives_restart(temp_db):
    """Verifies scheduler persists jobs in DB, survives restart, and executes due items."""
    scheduler1 = BackgroundScheduler(db_manager=temp_db)
    
    # 1. Add a reminder due immediately (delay_seconds=0)
    job = scheduler1.add_reminder(title="Submit Project Timesheet", delay_seconds=0)
    job_id = job["id"]
    assert job["enabled"] is True

    # 2. Simulate complete restart: instantiate new scheduler from same DB
    scheduler2 = BackgroundScheduler(db_manager=temp_db)
    jobs = scheduler2.list_jobs()
    assert any(j["id"] == job_id for j in jobs)

    # 3. Setup execution callback tracking
    triggered_events = []
    def on_reminder(evt):
        triggered_events.append(evt)
    scheduler2.add_listener(on_reminder)

    # 4. Execute due jobs
    executed_count = await scheduler2.execute_due_jobs()
    assert executed_count >= 1
    assert len(triggered_events) == 1
    assert triggered_events[0]["title"] == "Submit Project Timesheet"

    # 5. Check run log recorded
    logs = scheduler2.get_run_logs()
    assert len(logs) >= 1
    assert logs[0]["job_id"] == job_id
    assert logs[0]["status"] == "success"

    # 6. Job should now be disabled (one-time reminder)
    active_jobs = scheduler2.list_jobs(include_disabled=False)
    assert not any(j["id"] == job_id for j in active_jobs)

@pytest.mark.asyncio
async def test_live_data_fallback_to_scraping_and_labelled():
    """Verifies that API failure falls back to web scraping with source='scraped'."""
    # Force scrape mode to simulate API unavailability
    pipeline = LiveDataPipeline(force_scrape_for_testing=True)

    # Weather
    wx = await pipeline.get_weather("London")
    assert wx["source"] == "scraped"
    assert "location" in wx

    # Mutual fund
    mf = await pipeline.get_mutual_fund("Axis Bluechip")
    assert mf["source"] == "scraped"
    assert mf["nav"] > 0

    # Sports
    sports = await pipeline.get_sports_score("football", "Arsenal vs Chelsea")
    assert sports["source"] == "scraped"

@pytest.mark.asyncio
async def test_financial_watchlist_alert_triggers_once(temp_db):
    """Verifies watchlist price alerts trigger exactly once when condition is met."""
    fin = FinancialEngine(db_manager=temp_db)

    # Add alert: AAPL above $180
    fin.add_to_watchlist("AAPL", target_price=180.0, condition="above", asset_type="stock")

    # 1. Evaluate with price below target ($175) -> no alert
    alerts1 = await fin.check_alerts(quotes_override={"AAPL": 175.0})
    assert len(alerts1) == 0

    # 2. Evaluate with price above target ($185) -> triggers alert once
    alerts2 = await fin.check_alerts(quotes_override={"AAPL": 185.0})
    assert len(alerts2) == 1
    assert alerts2[0]["symbol"] == "AAPL"
    assert alerts2[0]["current_price"] == 185.0

    # 3. Evaluate again with higher price ($190) -> MUST NOT trigger duplicate alert
    alerts3 = await fin.check_alerts(quotes_override={"AAPL": 190.0})
    assert len(alerts3) == 0

    # Verify state in DB
    items = fin.list_watchlist()
    assert items[0]["triggered"] is True

@pytest.mark.asyncio
async def test_hard_prohibitions_enforced_in_code():
    """Verifies code-level prohibitions on email sending, calendar edit/delete, and financial trading."""
    gw = GoogleWorkspaceClient()
    fin = FinancialEngine()

    # 1. Sending Email must raise EmailSendingProhibitedError
    with pytest.raises(EmailSendingProhibitedError):
        gw.send_email("boss@example.com", "Project Status", "Everything is done.")

    # 2. Deleting calendar event must raise CalendarModificationProhibitedError
    with pytest.raises(CalendarModificationProhibitedError):
        gw.delete_calendar_event("event_123")

    # 3. Editing calendar event must raise CalendarModificationProhibitedError
    with pytest.raises(CalendarModificationProhibitedError):
        gw.edit_calendar_event("event_123", {"summary": "Changed Time"})

    # 4. Trading must raise TradingProhibitedError
    with pytest.raises(TradingProhibitedError):
        fin.place_trade("AAPL", "BUY", 10.0)

    # 5. Financial advice/suggestions must raise TradingProhibitedError
    with pytest.raises(TradingProhibitedError):
        fin.suggest_trade("BTC")

    # 6. Confirm Safety Layer has all prohibited tools registered as Tier 3
    prohibited_tools = ["send_email", "calendar_delete_event", "calendar_edit_event", "place_trade", "suggest_trade"]
    for t_name in prohibited_tools:
        tool_def = registry.get(t_name)
        assert tool_def is not None, f"Tool {t_name} must be registered"
        assert tool_def.tier == 3, f"Tool {t_name} must be Tier 3 blocked"

def test_each_document_type_opens_correctly(tmp_path):
    """Verifies creation and parsing of Word, Excel, PowerPoint, and PDF documents."""
    generator = DocumentGenerator(default_dir=str(tmp_path))

    # 1. Word Document (.docx)
    word_res = generator.create_word_document(
        filename="TestReport.docx",
        title="Phenom Phase 5 Architecture",
        sections=[
            {"heading": "Executive Summary", "content": "Assistant operating modes and integrations."},
            {"heading": "Key Milestones", "bullets": ["Scheduler Engine", "Document Generation", "Workspace Client"]}
        ]
    )
    docx_file = word_res["file_path"]
    assert os.path.exists(docx_file)
    # Open and verify with python-docx
    doc = docx.Document(docx_file)
    headings = [p.text for p in doc.paragraphs if p.text]
    assert any("Phenom Phase 5 Architecture" in h for h in headings)
    assert any("Executive Summary" in h for h in headings)

    # 2. Excel Spreadsheet (.xlsx)
    excel_res = generator.create_excel_spreadsheet(
        filename="FinancialSummary.xlsx",
        sheet_title="Portfolio Overview",
        headers=["Symbol", "Asset Type", "Current Price", "Target Alert"],
        rows=[
            ["AAPL", "Stock", 185.50, 180.00],
            ["BTC", "Crypto", 65000.00, 70000.00],
            ["NVDA", "Stock", 130.25, 125.00]
        ]
    )
    xlsx_file = excel_res["file_path"]
    assert os.path.exists(xlsx_file)
    # Open and verify with openpyxl
    wb = openpyxl.load_workbook(xlsx_file)
    ws = wb.active
    assert ws.title == "Portfolio Overview"
    assert ws.cell(row=1, column=1).value == "Symbol"
    assert ws.cell(row=2, column=1).value == "AAPL"
    assert ws.cell(row=3, column=1).value == "BTC"

    # 3. PowerPoint Presentation (.pptx)
    ppt_res = generator.create_powerpoint_presentation(
        filename="AssistantBriefing.pptx",
        title="Phenom Assistant System",
        subtitle="Desktop AI for Windows 11",
        slides_data=[
            {"title": "Three Assistant Modes", "bullets": ["Work: Email drafts & Calendar", "Study: Research & Vault", "Chill: Banter & Low Resource"]},
            {"title": "Safety Guarantees", "bullets": ["No email sending", "No financial trading", "100% Reversible cleanup"]}
        ]
    )
    pptx_file = ppt_res["file_path"]
    assert os.path.exists(pptx_file)
    # Open and verify with python-pptx
    prs = pptx.Presentation(pptx_file)
    assert len(prs.slides) == 3
    slide1_text = [shape.text for shape in prs.slides[0].shapes if shape.has_text_frame]
    assert any("Phenom Assistant System" in t for t in slide1_text)

    # 4. PDF Document (.pdf)
    pdf_res = generator.create_pdf_report(
        filename="SystemAudit.pdf",
        title="Phenom Security Audit",
        paragraphs=[
            "This report documents Phase 5 compliance with all permanent safety rules.",
            "All prohibited actions are blocked at the code level and registered as Tier 3."
        ],
        table_data=[
            ["Capability", "Risk Tier", "Enforcement"],
            ["Send Email", "Tier 3", "Blocked in Code"],
            ["Execute Trade", "Tier 3", "Blocked in Code"],
            ["Save Draft", "Tier 1", "Confirmation Dialog"],
            ["Create Event", "Tier 2", "Previewed Confirmation"]
        ]
    )
    pdf_file = pdf_res["file_path"]
    assert os.path.exists(pdf_file)
    # Open and verify with pypdf
    reader = pypdf.PdfReader(pdf_file)
    assert len(reader.pages) >= 1
    extracted_text = reader.pages[0].extract_text()
    assert "Phenom Security Audit" in extracted_text
    assert "Send Email" in extracted_text

def test_modes_switching_and_tone_adaptation(temp_db):
    """Verifies Work, Study, and Chill modes alter tone, tools, and notification rules."""
    modes = ModeManager(db_manager=temp_db)

    # Initial default is Work
    assert modes.get_current_mode() == AssistantMode.WORK
    assert modes.is_notification_allowed("normal") is True

    # Switch to Study
    res_study = modes.set_mode("study")
    assert modes.get_current_mode() == AssistantMode.STUDY
    assert "STUDY" in modes.get_tone_prompt()
    # Study mode suppresses normal notifications (quiet)
    assert modes.is_notification_allowed("normal") is False
    assert modes.is_notification_allowed("urgent") is True

    # Switch to Chill
    res_chill = modes.set_mode("chill")
    assert modes.get_current_mode() == AssistantMode.CHILL
    assert "CHILL" in modes.get_tone_prompt()
    assert modes.get_mode_profile()["resource_profile"] == "low"

    # Verify persistence across restart
    modes2 = ModeManager(db_manager=temp_db)
    assert modes2.get_current_mode() == AssistantMode.CHILL

@pytest.mark.asyncio
async def test_tech_news_briefing_aggregation():
    """Verifies tech news aggregation returns structured articles."""
    briefing_tool = TechNewsBriefing()
    res = await briefing_tool.get_briefing(items_per_source=1)
    assert "articles" in res
    assert len(res["articles"]) > 0
    art = res["articles"][0]
    assert "title" in art
    assert "link" in art
    assert "source" in art

@pytest.mark.asyncio
async def test_google_workspace_draft_and_calendar():
    """Verifies allowed Google Workspace operations: drafts and calendar creation."""
    gw = GoogleWorkspaceClient()

    # Draft creation (allowed)
    draft = gw.create_draft("client@example.com", "Project Deliverables", "Attached are the docs.")
    assert draft["status"] == "DRAFT"
    assert draft["to"] == "client@example.com"
    assert len(gw.list_drafts()) >= 1

    # Calendar event creation (allowed)
    event = gw.create_event("Sprint Demo", "2026-10-30T10:00:00Z", "2026-10-30T11:00:00Z", "Demo of Phase 5")
    assert event["summary"] == "Sprint Demo"
    assert len(gw.list_events()) >= 3
