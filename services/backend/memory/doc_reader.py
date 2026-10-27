"""
Phenom Multi-Format Document Reader & Text Chunker
Supports: PDF (.pdf), Word (.docx), PowerPoint (.pptx), Excel (.xlsx), Markdown (.md), and Text (.txt).
Provides overlap chunking for vector embeddings.
"""

import os
from typing import List, Dict, Any, Optional

class DocumentReader:
    @staticmethod
    def read_document(file_path: str) -> str:
        """Reads text from supported document types."""
        if not os.path.isfile(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()
        
        if ext in (".txt", ".md", ".json", ".log"):
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                return f.read()

        elif ext == ".pdf":
            import pypdf
            reader = pypdf.PdfReader(file_path)
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n\n".join(pages)

        elif ext == ".docx":
            import docx
            doc = docx.Document(file_path)
            paras = [p.text for p in doc.paragraphs if p.text]
            return "\n".join(paras)

        elif ext == ".pptx":
            import pptx
            prs = pptx.Presentation(file_path)
            slide_texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_texts.append(shape.text)
            return "\n".join(slide_texts)

        elif ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            cells = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(v) for v in row if v is not None]
                    if row_vals:
                        cells.append(" | ".join(row_vals))
            return "\n".join(cells)

        else:
            raise ValueError(f"Unsupported file format: {ext}")

    @staticmethod
    def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
        """Splits long text into overlapping chunks."""
        words = text.split()
        if not words:
            return []
        
        chunks = []
        step = max(1, chunk_size - overlap)
        for i in range(0, len(words), step):
            chunk = " ".join(words[i : i + chunk_size])
            chunks.append(chunk)
            if i + chunk_size >= len(words):
                break
        return chunks

doc_reader = DocumentReader()
